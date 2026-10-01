#!/usr/bin/env python3
"""Convert exportTables.sh UTF-8 TSV output using the original report template.

Environment setup and conversion: bash exportXlsx.sh INPUT.utf8.txt OUTPUT.xlsx
The wrapper installs requirements-xlsx.txt in a dedicated virtual environment.
Usage: python3 tsv_to_xlsx.py INPUT.utf8.txt OUTPUT.xlsx [--template TEMPLATE.xlsx]

Preserves the template SQL sheet, conditional formatting, and column styles.
The default template is resolved relative to this script, not the working directory.
"""

import argparse
from copy import copy
import re
from pathlib import Path
from tempfile import NamedTemporaryFile

from openpyxl import load_workbook

DEFAULT_TEMPLATE = (
    Path(__file__).resolve().parent
    / "report.special_case_taxa.template.xlsx"
)

# MariaDB optionally quotes fields and escapes quotes/backslashes with backslashes.
FIELD = re.compile(r'("(?:[^"\\]|\\.)*"|(?:[^\t\\]|\\.)*)(\t|$)')
ESCAPES = {"0": "\0", "b": "\b", "n": "\n", "r": "\r", "t": "\t", "Z": "\x1a"}


def parse_row(line):
    """Decode one server-exported record, preserving NULL versus literal text."""
    line = line.removesuffix("\n")
    position = 0
    while True:
        match = FIELD.match(line, position)
        if match is None:
            raise ValueError(f"Invalid TSV field at position {position}")
        raw, separator = match.groups()
        if raw == r"\N":
            value = None
        else:
            if raw.startswith('"') and raw.endswith('"'):
                raw = raw[1:-1]
            value = re.sub(r"\\(.)", lambda m: ESCAPES.get(m[1], m[1]), raw)
        yield value
        if not separator:
            break
        position = match.end()


def convert(source_path, destination_path, template_path=DEFAULT_TEMPLATE):
    source_path, destination_path, template_path = map(
        Path, (source_path, destination_path, template_path)
    )
    if destination_path.resolve() in {source_path.resolve(), template_path.resolve()}:
        raise ValueError("Output must differ from the input and template files")
    workbook = load_workbook(template_path)
    sheet = workbook["taxa_of_interest"]
    template_columns = [cell.value for cell in sheet[1]]
    # Reuse the first data row's per-column styles, including hyperlink styling.
    styles = [copy(cell._style) for cell in sheet[2]]
    numeric_columns = {
        index for index, cell in enumerate(sheet[2], start=1)
        if isinstance(cell.value, int) and not isinstance(cell.value, bool)
    }

    with source_path.open(encoding="utf-8", newline="") as source:
        header = source.readline()
        if not header:
            raise ValueError("Input file is empty")
        # exportTables.sh writes the unquoted client header separately.
        columns = header.rstrip("\r\n").split("\t")
        if columns != template_columns:
            raise ValueError("TSV columns must match the template headers and order")
        # Clear all old records so shorter exports cannot retain stale data.
        sheet.delete_rows(2, sheet.max_row - 1)

        for row_number, line in enumerate(source, start=2):
            values = list(parse_row(line))
            if len(values) != len(columns):
                raise ValueError(
                    f"Row {row_number}: expected {len(columns)} columns, got {len(values)}"
                )
            for column, value in enumerate(values, start=1):
                if column in numeric_columns and value is not None:
                    # Match the source workbook's numeric release and ID columns.
                    # Larger integers stay text to avoid Excel precision loss.
                    if re.fullmatch(r"-?\d{1,15}", value):
                        value = int(value)
                cell = sheet.cell(row_number, column, value)
                cell._style = copy(styles[column - 1])
                if isinstance(value, str):
                    # Keep IDs and other text unchanged. Only the report's URL
                    # columns may activate its existing Excel hyperlink formulas.
                    hyperlink = columns[column - 1].endswith("_url") and value.startswith('=HYPERLINK("')
                    cell.data_type = "f" if hyperlink else "s"

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    # Publish only a complete workbook; preserve an older export on save failure.
    temporary_path = None
    try:
        with NamedTemporaryFile(
            dir=destination_path.parent, prefix=f".{destination_path.name}.",
            suffix=".xlsx", delete=False,
        ) as temporary:
            temporary_path = Path(temporary.name)
        workbook.save(temporary_path)
        temporary_path.replace(destination_path)
    finally:
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)
    return sheet.max_row - 1


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    args = parser.parse_args()
    if args.output.resolve() in {args.input.resolve(), args.template.resolve()}:
        parser.error("Output must differ from the input and template files")
    rows = convert(args.input, args.output, args.template)
    print(f"Saved {rows:,} rows to {args.output}")


if __name__ == "__main__":
    main()
